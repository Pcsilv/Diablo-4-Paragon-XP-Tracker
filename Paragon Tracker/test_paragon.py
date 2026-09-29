"""Testes automáticos:  python -m unittest -v test_paragon"""
import os
import random
import unittest

import openpyxl

import paragon_core as pc

XLSX = pc.find_xlsx()
T = pc.load_table(XLSX)

# Valor informado pelo usuário como referência do total (usado SOMENTE como expectativa de teste)
EXPECTED_TOTAL = 58_237_517_372


def raw_sheet_columns():
    """Leitura independente: colunas A, B, C direto da planilha (linhas 2..301)."""
    ws = openpyxl.load_workbook(XLSX, data_only=True).worksheets[0]
    rows = [(ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(2, 302)]
    return rows


class TestTabela(unittest.TestCase):
    def test_total_bate(self):
        self.assertEqual(T.max_level, 300)
        self.assertEqual(T.total, EXPECTED_TOTAL)

    def test_verificacao_independente(self):
        # 1) soma da coluna B lida por outro caminho
        rows = raw_sheet_columns()
        self.assertEqual(sum(b for _, b, _ in rows), EXPECTED_TOTAL)
        # 2) coluna "Cumulative XP" da planilha == nosso acumulado, nível a nível
        for lv, b, c in rows:
            self.assertEqual(T.cum[lv], c, f"acumulado difere no P{lv}")
            self.assertEqual(T.req[lv], b)
        self.assertEqual(T.warnings, [])

    def test_tudo_inteiro(self):
        self.assertTrue(all(type(x) is int for x in T.req[1:] + T.cum))


class TestNiveisChave(unittest.TestCase):
    def test_marcos(self):
        for L in (1, 100, 200, 250, 279, 280, 290, 299, 300):
            p = T.locate(T.cum[L])
            self.assertEqual((p.level, p.in_level), (L, 0), f"P{L}")
            if L < 300:
                self.assertEqual(p.span, T.req[L + 1])
                p2 = T.locate(T.cum[L] + T.span(L) - 1)
                self.assertEqual((p2.level, p2.in_level), (L, T.span(L) - 1))
                p3 = T.locate(T.cum[L] + T.span(L))  # exatamente o necessário p/ passar
                self.assertEqual((p3.level, p3.in_level), (L + 1, 0))
        print("\n  Marcos (XP para alcançar o Paragon):")
        for L in (1, 100, 200, 250, 279, 280, 290, 299, 300):
            print(f"   P{L:<3} -> {pc.fmt_int(T.cum[L])}")

    def test_p250_exemplo(self):
        self.assertEqual(T.cum[250], 10_656_835_750)
        p = T.locate(T.cum[250] + 100_000_000)
        self.assertEqual(p.level, 250)
        self.assertEqual(p.in_level, 100_000_000)
        self.assertEqual(p.span, T.req[251])
        self.assertEqual(p.level_remaining, T.req[251] - 100_000_000)
        self.assertEqual(p.total_remaining, T.total - T.cum[250] - 100_000_000)
        self.assertEqual(p.next_level, 251)

    def test_todos_os_limites(self):
        for L in range(0, 301):
            lo = T.locate(T.cum[L])
            self.assertEqual((lo.level, lo.in_level), (L, 0))
            if L < 300:
                hi = T.locate(T.cum[L + 1] - 1)
                self.assertEqual((hi.level, hi.in_level), (L, T.span(L) - 1))

    def test_ida_e_volta_aleatoria(self):
        rnd = random.Random(1234)
        for _ in range(20000):
            xp = rnd.randint(0, T.total)
            p = T.locate(xp)
            xp2, w = T.resolve_level(p.level, p.in_level)
            self.assertEqual(xp2, xp)
            self.assertEqual(w, [])
            self.assertEqual(p.xp, xp)
            self.assertEqual(p.total_remaining, T.total - xp)

    def test_percentuais_nunca_decrescem(self):
        prev = -1
        for L in range(0, 301):
            bp = pc.pct_bp(T.cum[L], T.total)
            self.assertGreaterEqual(bp, prev)
            prev = bp


class TestLimites(unittest.TestCase):
    def test_zero(self):
        p = T.locate(0)
        self.assertEqual((p.level, p.in_level, p.span, p.total_remaining), (0, 0, T.req[1], T.total))
        self.assertEqual(pc.fmt_pct(p.xp, T.total), "0,00%")

    def test_p300_exato(self):
        p = T.locate(T.total)
        self.assertEqual(p.level, 300)
        self.assertEqual(p.total_remaining, 0)
        self.assertEqual(p.level_remaining, 0)
        self.assertIsNone(p.next_level)
        self.assertEqual(pc.fmt_pct(p.xp, T.total), "100,00%")

    def test_acima_do_maximo(self):
        xp, w = T.resolve_total(T.total + 1)
        self.assertEqual(xp, T.total)
        self.assertEqual(len(w), 1)
        self.assertEqual(T.locate(T.total * 3).xp, T.total)
        xp, w = T.resolve_level(300, 5)
        self.assertEqual(xp, T.total)
        self.assertTrue(w)
        xp, w = T.resolve_level(301, 0)
        self.assertEqual(xp, T.total)
        self.assertTrue(w)

    def test_xp_dentro_do_nivel_excessivo(self):
        xp, w = T.resolve_level(250, T.req[251])  # exatamente o necessário p/ passar => corrigido
        self.assertEqual(xp, T.cum[251] - 1)
        self.assertTrue(w)

    def test_negativos_e_texto(self):
        for bad in ("-1", "-100", "abc", "12a", "", "1,5", "10.5x"):
            with self.assertRaises(ValueError, msg=bad):
                pc.parse_int(bad)
        self.assertEqual(pc.parse_int("10.656.835.750"), 10_656_835_750)
        self.assertEqual(pc.parse_int(" 10656835750 "), 10_656_835_750)
        with self.assertRaises(ValueError):
            pc.parse_int("1.5", thousands=False)

    def test_paragon_fora_do_intervalo(self):
        xp, w = T.resolve_level(-3, 0)
        self.assertEqual(xp, 0)
        self.assertTrue(w)


class TestPercentuais(unittest.TestCase):
    def test_50_por_cento_exato(self):
        # metade do total exata (total é par)
        self.assertEqual(T.total % 2, 0)
        self.assertEqual(pc.fmt_pct(T.total // 2, T.total), "50,00%")
        self.assertEqual(pc.fmt_pct(29_118_758_686, T.total), "50,00%")

    def test_nao_mostra_100_antes_da_hora(self):
        self.assertEqual(pc.fmt_pct(T.total - 1, T.total), "99,99%")

    def test_formatacao(self):
        self.assertEqual(pc.fmt_int(58_237_517_372), "58.237.517.372")
        self.assertEqual(pc.fmt_int(0), "0")
        self.assertEqual(pc.fmt_bp(4837), "48,37%")


class TestPlanilhaSubstituivel(unittest.TestCase):
    def test_outra_curva(self):
        t = pc.ParagonTable([10, 20, 30])
        self.assertEqual((t.max_level, t.total, t.cum), (3, 60, [0, 10, 30, 60]))
        p = t.locate(35)
        self.assertEqual((p.level, p.in_level, p.span), (2, 5, 30))


if __name__ == "__main__":
    unittest.main(verbosity=2)
